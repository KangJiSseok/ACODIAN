"""Office 문서(docx/pptx/xlsx) 텍스트 + 임베디드 이미지 추출."""

from __future__ import annotations

import logging
from dataclasses import dataclass, field
from pathlib import Path

logger = logging.getLogger(__name__)

MAX_CHARS = 60_000
MAX_IMAGES = 20
MAX_IMAGE_BYTES = 5 * 1024 * 1024


@dataclass
class ExtractedImage:
    data: bytes
    mime_type: str
    origin: str


@dataclass
class OfficeContent:
    text: str
    images: list[ExtractedImage] = field(default_factory=list)


class OfficeExtractionError(RuntimeError):
    pass


def extract_office_content(path: Path) -> OfficeContent:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        content = _extract_docx(path)
    elif suffix == ".pptx":
        content = _extract_pptx(path)
    elif suffix == ".xlsx":
        content = _extract_xlsx(path)
    else:
        raise OfficeExtractionError(f"지원하지 않는 office 확장자: {suffix}")

    content.text = content.text.strip()
    if not content.text and not content.images:
        raise OfficeExtractionError(f"읽을 수 있는 내용이 없습니다: {path.name}")
    if len(content.text) > MAX_CHARS:
        logger.info("Office 텍스트 truncate: %s %d→%d", path.name, len(content.text), MAX_CHARS)
        content.text = content.text[:MAX_CHARS]
    content.images = _cap_images(path.name, content.images)
    return content


def _cap_images(file_name: str, images: list[ExtractedImage]) -> list[ExtractedImage]:
    kept: list[ExtractedImage] = []
    for image in images:
        if len(image.data) > MAX_IMAGE_BYTES:
            logger.info("oversized image 제외 (%d bytes) from %s", len(image.data), file_name)
            continue
        kept.append(image)
        if len(kept) >= MAX_IMAGES:
            logger.info("image cap (%d) 도달, 나머지 무시 from %s", MAX_IMAGES, file_name)
            break
    return kept


def _extract_docx(path: Path) -> OfficeContent:
    from docx import Document
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = Document(str(path))
    chunks: list[str] = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            chunks.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                chunks.append(" | ".join(cells))

    images: list[ExtractedImage] = []
    for rel in document.part.rels.values():
        if rel.reltype != RT.IMAGE:
            continue
        try:
            part = rel.target_part
            images.append(
                ExtractedImage(
                    data=part.blob,
                    mime_type=part.content_type,
                    origin="docx body",
                )
            )
        except Exception:
            logger.warning("docx 이미지 읽기 실패 %s in %s", rel.rId, path.name)

    return OfficeContent(text="\n".join(chunks), images=images)


def _extract_pptx(path: Path) -> OfficeContent:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(path))
    chunks: list[str] = []
    images: list[ExtractedImage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        chunks.append(f"## Slide {index}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if text:
                        chunks.append(text)
            elif getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        chunks.append(" | ".join(cells))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    picture = shape.image
                    images.append(
                        ExtractedImage(
                            data=picture.blob,
                            mime_type=picture.content_type,
                            origin=f"slide {index}",
                        )
                    )
                except Exception:
                    logger.warning("pptx slide %d 이미지 읽기 실패 in %s", index, path.name)
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        if notes:
            chunks.append(f"[notes] {notes}")
    return OfficeContent(text="\n".join(chunks), images=images)


def _extract_xlsx(path: Path) -> OfficeContent:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), data_only=True)
    chunks: list[str] = []
    images: list[ExtractedImage] = []
    try:
        for sheet in workbook.worksheets:
            chunks.append(f"## Sheet: {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if cell is None else str(cell) for cell in row]
                if any(cell.strip() for cell in cells):
                    chunks.append(" | ".join(cells))
            for image in getattr(sheet, "_images", []) or []:
                blob = _xlsx_image_bytes(image)
                if blob is None:
                    continue
                images.append(
                    ExtractedImage(
                        data=blob,
                        mime_type=_xlsx_image_mime(image),
                        origin=f"sheet {sheet.title}",
                    )
                )
    finally:
        workbook.close()
    return OfficeContent(text="\n".join(chunks), images=images)


def _xlsx_image_bytes(image: object) -> bytes | None:
    ref = getattr(image, "ref", None)
    if hasattr(ref, "getvalue"):
        return ref.getvalue()
    if isinstance(ref, (bytes, bytearray)):
        return bytes(ref)
    data_method = getattr(image, "_data", None)
    if callable(data_method):
        try:
            return data_method()
        except Exception:
            return None
    return None


def _xlsx_image_mime(image: object) -> str:
    fmt = (getattr(image, "format", "") or "").lower()
    if fmt in {"jpg", "jpeg"}:
        return "image/jpeg"
    if fmt == "gif":
        return "image/gif"
    if fmt == "bmp":
        return "image/bmp"
    return "image/png"
